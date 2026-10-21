# -*- coding: utf-8 -*-
'''
EXPOsan: Exposition of sanitation and resource recovery systems

This module is developed by:
    Ali Ahmad <aliahmad1331@gmail.com>

This module is under the University of Illinois/NCSA Open Source License.
Please refer to https://github.com/QSD-Group/EXPOsan/blob/main/LICENSE.txt
for license details.
'''
from pathlib import Path
import re
import math
import pandas as pd
import numpy as np
import qsdsan as qs
import biosteam as bst
from exposan.biobinder_ml._distill_utils_flex import get_distillation_trials
from ua_column_diagnostics import add_crudeheavy_dis_diagnostics, solve_IRR_aitken_limit


def collect(system, case, settings, simulation_error='', console_log=''):
    sys=system;tea=sys.TEA;lca=getattr(sys,'LCA',None);hours=sys.operating_hours
    blocks=[("Scenario settings",pd.DataFrame(list(settings.items()),columns=["Assumption","Value"]))];diagnostics=[];trials=[]
    # Snapshot immediately, including partial searches from failed simulations.
    for col in sys.units:
        if hasattr(col, 'LHK'):
            for attempt, record in enumerate(get_distillation_trials(col), start=1):
                trials.append(dict(record, System=case, Attempt=attempt))
    def block(title,calculation):
        try:
            value=calculation()
            frame=value if isinstance(value,pd.DataFrame) else pd.DataFrame(value)
        except Exception as e:
            frame=pd.DataFrame([{'Unavailable':f'{type(e).__name__}: {e}'}])
        blocks.append((title,frame))
    def rows(items):return pd.DataFrame(items,columns=['Item','Value'])
    # Existing UA IRR reporter; no new root/finance implementation.
    irr,reason=(np.nan,'simulation_failed')
    if not simulation_error:
        try:irr,reason=solve_IRR_aitken_limit(tea)
        except Exception as e:reason=f'{type(e).__name__}: {e}'
    metrics={'Status':'FAIL' if simulation_error else 'PASS',
             'IRR_pct':float(irr)*100 if np.isfinite(irr) else None,
             'IRR_reason':reason,'NPV':None if simulation_error else tea.NPV,'Operating_hours':hours,
             'Cutoff_fracs':settings.get('Native_cut_fractions'),
             'Intermediate_route':settings.get('Intermediate_route'),
             'Error':simulation_error}
    # Native totals named in the uploaded TEA sheet.
    for name in ('TDC','FCI','TCI','land','AOC','FOC','VOC',
                 'ISBL_installed_equipment_cost','OSBL_installed_equipment_cost'):
        try:
            value=getattr(tea,name)
            metrics['tea.'+name]=value() if callable(value) else value
        except Exception:metrics['tea.'+name]=None
    for name in ('material_cost','sales','utility_cost'):
        try:metrics['sys.'+name]=getattr(sys,name) if not simulation_error else None
        except Exception:metrics['sys.'+name]=None
    block('TEA TOTALS',lambda:rows([(k,v) for k,v in metrics.items()]))
    # Exact selection condition and cost*hours from workbook snippets.
    block('FEEDS — sys.material_cost [USD/year]',lambda:rows([
        (i.ID,i.cost*hours) for i in sys.feeds
        if i.cost or 'surrogate' in i.ID or 'recovered' in i.ID]))
    block('PRODUCTS — sys.sales [USD/year]',lambda:rows([
        (i.ID,i.cost*hours) for i in sys.products
        if i.cost or 'surrogate' in i.ID or 'recovered' in i.ID]))
    block('ISBL INSTALLED COST [USD]',lambda:rows([
        (i.ID,i.installed_cost) for i in sys.units if i not in tea.OSBL_units
        and (i.installed_cost or 'EC' in i.ID)]))
    block('OSBL INSTALLED COST [USD]',lambda:rows([
        (i.ID,i.installed_cost) for i in sys.units if i in tea.OSBL_units
        and (i.installed_cost or 'CHP' in i.ID)]))
    block('UNIT UTILITIES [USD/year]',lambda:rows([
        (i.ID,i.utility_cost*hours) for i in sys.units if i.utility_cost or i.ID=='EC']))
    block('UNIT ELECTRICITY [USD/year]',lambda:rows([
        (i.ID,i.power_utility.rate*qs.PowerUtility.price*hours) for i in sys.units
        if i.power_utility.rate or i.ID=='EC']))
    for agent_ID in ('high_pressure_steam','ethylene','cooling_water','chilled_water','natural_gas'):
        def heat_rows(agent_ID=agent_ID):
            agent=qs.HeatUtility.get_agent(agent_ID)
            return rows([(i.ID,hu.cost*hours) for i in sys.units for hu in i.heat_utilities
                         if hu and hu.agent.ID==agent.ID])
        block(agent_ID.upper()+' BY UNIT [USD/year]',heat_rows)
    block('COOLING BY UNIT [USD/year]',lambda:rows([
        (i.ID,sum(hu.cost for hu in i.heat_utilities if hu.ID in ('cooling_water','chilled_water'))*hours)
        for i in sys.units if any(hu.cost for hu in i.heat_utilities if hu.ID in ('cooling_water','chilled_water'))]))
    block('CHP SYSTEM HEATING UTILITIES [USD/year]',lambda:rows([
        (k,sum(i.cost for i in v)*hours) for k,v in sys.flowsheet.unit.CHP.sys_heating_utilities.items()]))
    def stream_utilities(side):
        values=[];prices=bst.stream_utility_prices
        for i in sys.units:
            indices=getattr(i,'_'+side+'_utility_indices')
            for k,v in indices.items():
                stream=(i.ins if side=='inlet' else i.outs)[v]
                if stream.F_mass and not stream.price:
                    values.append((i.ID+' / '+k,stream.F_mass*prices[k]*hours))
        return rows(values)
    block('INLET STREAM UTILITIES [USD/year]',lambda:stream_utilities('inlet'))
    block('OUTLET STREAM UTILITIES [USD/year]',lambda:stream_utilities('outlet'))
    block('CASHFLOW TABLE',lambda:tea.get_cashflow_table().reset_index())
    # Exact native tables from LCA text boxes; no extra impact methods/allocation.
    for category,annual in [('stream',True),('Construction',False),('Transportation',False),('other',True)]:
        def impact(category=category,annual=annual):
            table=lca.get_impact_table(category,annual=True) if annual else lca.get_impact_table(category)
            return table.reset_index()
        block('LCA '+category.upper()+(' — ANNUAL' if annual else ' — NATIVE LIFETIME BASIS'),impact)
    # Existing UA diagnostic fields only, applied to each distillation column.
    for col in sys.units:
        if hasattr(col,'LHK') and len(col.outs)==2:
            d={'System':case,'Column':col.ID}
            try:add_crudeheavy_dis_diagnostics(d,col,prefix='')
            except Exception as e:d['Diagnostic_error']=f'{type(e).__name__}: {e}'
            diagnostics.append(d)
    return dict(case=case,metrics=metrics,blocks=blocks,columns=diagnostics,trials=trials)


def failed_snapshot(case,settings,error,log):
    return dict(case=case,metrics={'Status':'FAIL','Error':error,'IRR_pct':None},
                blocks=[('ERROR',pd.DataFrame([{'Error':error}]))],columns=[],trials=[])


def export(snapshots,path,notes=None):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    def value(x):
        if isinstance(x,(np.integer,np.floating)):x=x.item()
        if isinstance(x,float) and not math.isfinite(x):return None
        if x is None or isinstance(x,(int,float,bool)):return x
        return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]','',str(x))
    def clean(df):
        df=df.copy();df.columns=[str(c) for c in df.columns]
        for c in df.columns:df[c]=df[c].map(value)
        return df
    with pd.ExcelWriter(path,engine='openpyxl') as writer:
        summary=pd.DataFrame({s['case']:s['metrics'] for s in snapshots}).rename_axis('Metric').reset_index()
        clean(summary).to_excel(writer,sheet_name='Comparison',index=False)
        diagnosis=[];trial_records=[]
        for s in snapshots:
            start=0
            for title,table in s['blocks']:
                pd.DataFrame([[title]]).to_excel(writer,sheet_name=s['case'][:31],startrow=start,index=False,header=False)
                start+=1
                if not table.empty:
                    clean(table).to_excel(writer,sheet_name=s['case'][:31],startrow=start,index=False)
                    start+=len(table)+1
                start+=2
            diagnosis.extend(s['columns'])
            trial_records.extend(s.get('trials', []))
        clean(pd.DataFrame(diagnosis) if diagnosis else pd.DataFrame(columns=['System','Column'])).to_excel(writer,sheet_name='Distillation',index=False)
        trial_table = pd.DataFrame(trial_records)
        if trial_table.empty:
            trial_table = pd.DataFrame(columns=['System', 'Column', 'Attempt', 'Status', 'Reason'])
        else:
            leading = [c for c in ['System', 'Column', 'Attempt', 'Timestamp',
                'Status', 'Stage', 'Reason', 'LHK_light', 'LHK_heavy', 'Lr', 'Hr',
                'Top_ratio', 'Target_ratio', 'Tolerance', 'Target_error',
                'Purchase_cost', 'Stages', 'Reflux'] if c in trial_table]
            trial_table = trial_table[leading + [c for c in trial_table if c not in leading]]
        clean(trial_table).to_excel(writer, sheet_name='Distillation_trials', index=False)
        from openpyxl.styles import Font,PatternFill
        for ws in writer.book:
            ws.freeze_panes='B2' if ws.title=='Comparison' else 'A2'
            ws.sheet_view.showGridLines=False
            if ws.title=='Distillation_trials':ws.auto_filter.ref=ws.dimensions
            for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D')
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width=min(48,max(16,max(len(str(c.value or '')) for c in col[:50])+2))
            for row in ws:
                for c in row:
                    if isinstance(c.value,float):c.number_format='#,##0.00;[Red](#,##0.00)'
                    if c.data_type=='s' and str(c.value).startswith(('=','+','@')):c.data_type='s'
    return path
